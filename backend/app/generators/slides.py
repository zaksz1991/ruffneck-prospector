"""
Simple PowerPoint generator using python-pptx (free).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from typing import Dict, Any
import os
import uuid


def create_slides_pptx(
    topic: str,
    branding: Dict[str, Any],
    prospect: Dict[str, Any],
    output_dir: str = "/tmp",
) -> str:
    company = branding.get("company_name", "RuffNeck Entertainment")
    username = prospect.get("username") or "Reader"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1 - Title
    slide_layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(slide_layout)
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = topic
    p.font.size = Pt(36)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER

    sub = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
    tf2 = sub.text_frame
    p2 = tf2.paragraphs[0]
    p2.text = f"Prepared for @{username}  •  {company}"
    p2.font.size = Pt(16)
    p2.alignment = PP_ALIGN.CENTER

    # Slide 2 - Key points
    slide2 = prs.slides.add_slide(slide_layout)
    title2 = slide2.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11), Inches(0.8))
    title2.text_frame.paragraphs[0].text = "Key Takeaways"
    title2.text_frame.paragraphs[0].font.size = Pt(28)
    title2.text_frame.paragraphs[0].font.bold = True

    body = slide2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11), Inches(4.5))
    tf = body.text_frame
    tf.word_wrap = True
    points = [
        "Clarify your unique value in one sentence.",
        "Find conversations where your ideal audience is already active.",
        "Give first — share a useful insight before asking for anything.",
        "Track what resonates and double down.",
    ]
    for i, point in enumerate(points):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = f"→  {point}"
        p.font.size = Pt(18)
        p.space_after = Pt(14)

    # Slide 3 - CTA
    slide3 = prs.slides.add_slide(slide_layout)
    cta = slide3.shapes.add_textbox(Inches(1), Inches(2.8), Inches(11), Inches(1.5))
    p = cta.text_frame.paragraphs[0]
    p.text = "Ready for the next step?"
    p.font.size = Pt(28)
    p.alignment = PP_ALIGN.CENTER
    p.font.bold = True

    sub3 = slide3.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11), Inches(1))
    p3 = sub3.text_frame.paragraphs[0]
    p3.text = "Reply to the outreach message and tell us which takeaway you want help with."
    p3.font.size = Pt(16)
    p3.alignment = PP_ALIGN.CENTER

    filename = f"slides_{username}_{uuid.uuid4().hex[:8]}.pptx"
    filepath = os.path.join(output_dir, filename)
    prs.save(filepath)
    return filepath
